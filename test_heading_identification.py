"""测试 AI 标题识别功能。"""

import asyncio
import sys
from pathlib import Path

# 添加 backend 到路径
sys.path.insert(0, str(Path(__file__).parent / "backend"))

from docx import Document
from app.services.document_extractor import DocumentExtractor

# 创建一个测试文档
TEST_DOCX_PATH = Path("/tmp/test_no_heading.docx")

def create_test_docx():
    """创建一个没有标题样式的测试文档。"""
    doc = Document()

    # 添加封面
    doc.add_paragraph("公司名称：测试公司", style="Normal")
    doc.add_paragraph("项目名称：测试项目", style="Normal")
    doc.add_paragraph("日期：2024-01-01", style="Normal")
    doc.add_page_break()

    # 添加目录
    doc.add_paragraph("目录", style="Normal")
    doc.add_page_break()

    # 添加正文内容 - 完全没有标题特征的内容
    doc.add_paragraph("项目概述", style="Normal")  # 没有编号，纯文本
    doc.add_paragraph("本项目旨在开发一个智能文档排版工具，帮助用户快速将格式混乱的文档转换为标准格式。")
    doc.add_paragraph("项目背景", style="Normal")  # 没有编号
    doc.add_paragraph("随着人工智能的发展，文档处理需求日益增长。传统文档排版需要大量人工操作，效率低下。")
    doc.add_paragraph("项目目标", style="Normal")  # 没有编号
    doc.add_paragraph("提高文档排版效率，减少人工操作。通过 AI 技术自动识别文档结构和样式。")
    doc.add_paragraph("技术方案", style="Normal")  # 没有编号
    doc.add_paragraph("架构设计", style="Normal")  # 没有编号
    doc.add_paragraph("采用前后端分离架构，前端使用 React，后端使用 Python FastAPI。")
    doc.add_paragraph("技术栈选型", style="Normal")  # 没有编号
    doc.add_paragraph("前端采用 React 18 和 TypeScript，配合 Ant Design 组件库。")
    doc.add_paragraph("后端采用 Python 3.11 和 FastAPI 框架，数据库使用 PostgreSQL。")

    doc.save(TEST_DOCX_PATH)
    print(f"测试文档已创建：{TEST_DOCX_PATH}")

def test_extractor():
    """测试 DocumentExtractor 的标题识别。"""
    print("\n=== 测试 DocumentExtractor ===")

    extractor = DocumentExtractor(str(TEST_DOCX_PATH))
    sections = extractor.extract_all()

    print(f"提取的 sections 数量：{len(sections)}")

    heading_count = sum(1 for s in sections if s.get("paragraph_type") == "heading")
    body_count = sum(1 for s in sections if s.get("paragraph_type") == "body")

    print(f"标题数量：{heading_count}")
    print(f"正文数量：{body_count}")

    # 检查是否需要 AI 识别
    needs_ai = DocumentExtractor.has_insufficient_headings(sections)
    print(f"需要 AI 识别：{needs_ai}")

    # 显示 format hints
    _, format_text = extractor.extract_all_with_format_hints()
    print("\n=== 格式提示文本 ===")
    print(format_text[:500] + "..." if len(format_text) > 500 else format_text)

    return sections

if __name__ == "__main__":
    create_test_docx()
    test_extractor()
